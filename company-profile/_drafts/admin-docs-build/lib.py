# -*- coding: utf-8 -*-
"""Shared Word builder for SYLTRA SMART admin documents (Arabic, RTL, A4)."""
import re, os
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
WM_DARK = os.path.join(HERE, "wordmark-dark.png")
WM_LIGHT = os.path.join(HERE, "wordmark-light.png")

SANS = "IBM Plex Sans Arabic"; MONO = "IBM Plex Mono"
VOID, GRAPHITE, SLATE, CHROME, PLATINUM = "0B0C0E", "17181C", "5B6068", "C7CCD3", "ECEDEF"
ION, ION_PRINT, CARD, LINE = "4C8DFF", "2F6BE0", "F3F4F6", "D9DCE1"
AR_RE = re.compile(r"[؀-ۿ]")
PAGE_W = Mm(210); MARG = Mm(18); CONTENT_W = PAGE_W - 2 * MARG

COMPANY = "سيلترا سمارت"
COMPANY_EN = "SYLTRA SMART"
ADDRESS = "الرياض — حي الملقا"
PHONE = "+966 55 009 8550"
EMAIL = "info@syltraone.com"
WEB = "syltraone.com"

def rgb(h): return RGBColor.from_string(h)

def set_run_font(run, family, size, color=VOID, bold=False, rtl=None):
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts"); rPr.insert(0, rFonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(a), family)
    run.font.size = Pt(size)
    szCs = OxmlElement("w:szCs"); szCs.set(qn("w:val"), str(int(size * 2))); rPr.append(szCs)
    run.font.color.rgb = rgb(color)
    if bold:
        run.font.bold = True
        rPr.append(OxmlElement("w:bCs"))
    if rtl is None:
        rtl = bool(AR_RE.search(run.text))
    if rtl:
        r = OxmlElement("w:rtl"); r.set(qn("w:val"), "1"); rPr.append(r)

def set_par(p, align=None, before=0, after=0, line=1.0, keep_next=False, bidi=True):
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = line
    if keep_next: pf.keep_with_next = True
    pPr = p._p.get_or_add_pPr()
    b = OxmlElement("w:bidi"); b.set(qn("w:val"), "1" if bidi else "0"); pPr.append(b)
    if align is None: align = "right" if bidi else "left"
    if bidi:
        align = {"right": "start", "left": "end"}.get(align, align)
    jc = pPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc"); pPr.append(jc)
    jc.set(qn("w:val"), align)

def add_text(container, text, size=10.5, color=VOID, bold=False, align=None,
             before=0, after=4, line=1.45, keep_next=False, mono=False, family=None):
    p = container.add_paragraph()
    latin_only = not AR_RE.search(text)
    fam = family or (MONO if (mono and latin_only) else SANS)
    bidi = not latin_only
    set_par(p, align=align, before=before, after=after, line=line, keep_next=keep_next, bidi=bidi)
    r = p.add_run(text)
    set_run_font(r, fam, size, color, bold)
    return p

def add_runs(container, parts, size=10.5, align=None, before=0, after=4, line=1.45, keep_next=False):
    """parts: list of (text, bold) or (text, bold, color)."""
    p = container.add_paragraph()
    set_par(p, align=align, before=before, after=after, line=line, keep_next=keep_next, bidi=True)
    for part in parts:
        text, bold = part[0], part[1]
        color = part[2] if len(part) > 2 else VOID
        r = p.add_run(text); set_run_font(r, SANS, size, color, bold)
    return p

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("bottom", bottom), ("start", left), ("end", right)):
        e = OxmlElement("w:" + k); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
    tcPr.append(mar)

def cell_borders(cell, color=LINE, sz=6, sides=("top", "bottom", "start", "end")):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for s in ("top", "bottom", "start", "end"):
        e = OxmlElement("w:" + s)
        if s in sides and sz:
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), color)
        else:
            e.set(qn("w:val"), "nil")
        b.append(e)
    tcPr.append(b)

def row_height(row, mm, exact=False):
    row.height = Mm(mm)
    trPr = row._tr.get_or_add_trPr(); h = trPr.find(qn("w:trHeight"))
    h.set(qn("w:hRule"), "exact" if exact else "atLeast")

def keep_table(t):
    rows = t.rows
    for i, row in enumerate(rows):
        trPr = row._tr.get_or_add_trPr(); trPr.append(OxmlElement("w:cantSplit"))
        if i < len(rows) - 1:
            for c in row.cells:
                for p in c.paragraphs: p.paragraph_format.keep_with_next = True

def table_setup(t, widths, bidi=True):
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblPr = t._tbl.tblPr
    if bidi:
        tblPr.append(OxmlElement("w:bidiVisual"))
    for old in tblPr.findall(qn("w:tblLayout")): tblPr.remove(old)
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    st = tblPr.find(qn("w:tblStyle"))
    if st is not None: tblPr.remove(st)
    for row in t.rows:
        for i, c in enumerate(row.cells):
            c.width = widths[i]
    grid = t._tbl.find(qn("w:tblGrid"))
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[i] / 635)))

def clear_cell(cell):
    for p in cell.paragraphs:
        p._p.getparent().remove(p._p)

def spacer(container, pts):
    p = container.add_paragraph(); set_par(p, after=0, before=0, line=1.0)
    r = p.add_run(""); set_run_font(r, SANS, pts)

def page_break(doc):
    p = doc.add_paragraph(); set_par(p)
    p.add_run().add_break(WD_BREAK.PAGE)

# ---------- document skeleton ----------
def new_doc(title, subtitle, doc_code, header_title=None):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = PAGE_W; sec.page_height = Mm(297)
    sec.left_margin = sec.right_margin = MARG
    sec.top_margin = Mm(30); sec.bottom_margin = Mm(20)
    sec.header_distance = Mm(9); sec.footer_distance = Mm(8)
    st = doc.styles["Normal"]; st.font.name = SANS; st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:cs"), SANS)
    # header: wordmark right, doc code left
    hp = sec.header.paragraphs[0]
    ht = sec.header.add_table(rows=1, cols=2, width=CONTENT_W)
    table_setup(ht, [int(CONTENT_W * 0.5), int(CONTENT_W * 0.5)], bidi=False)
    for c in ht.rows[0].cells:
        cell_borders(c, LINE, 0); cell_margins(c, 0, 0, 0, 0); clear_cell(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    c0 = ht.rows[0].cells[0]
    p = c0.add_paragraph(); set_par(p, align="left", bidi=False, line=1.2)
    r = p.add_run(COMPANY_EN + " · " + doc_code); set_run_font(r, MONO, 7.5, SLATE, rtl=False)
    p = c0.add_paragraph(); set_par(p, align="left", bidi=False, line=1.2)
    r = p.add_run(header_title or title); set_run_font(r, SANS, 8, SLATE)
    c1 = ht.rows[0].cells[1]
    p = c1.add_paragraph(); set_par(p, align="right", bidi=False)
    p.add_run().add_picture(WM_DARK, width=Mm(26))
    hp._p.getparent().remove(hp._p)
    # rule under header
    rp = sec.header.add_paragraph(); set_par(rp, after=0, before=2, bidi=False)
    pPr = rp._p.get_or_add_pPr(); pb = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "single"); bt.set(qn("w:sz"), "8"); bt.set(qn("w:color"), ION_PRINT); pb.append(bt); pPr.append(pb)
    r = rp.add_run(""); set_run_font(r, SANS, 2)
    # footer
    fp = sec.footer.paragraphs[0]; set_par(fp, align="center", bidi=True, line=1.2)
    r = fp.add_run(COMPANY + " · " + ADDRESS + " · "); set_run_font(r, SANS, 7.5, SLATE)
    r = fp.add_run(PHONE + " · " + EMAIL + " · " + WEB); set_run_font(r, MONO, 7.5, SLATE, rtl=False)
    fp2 = sec.footer.add_paragraph(); set_par(fp2, align="center", bidi=True, line=1.2)
    r = fp2.add_run("صفحة "); set_run_font(r, SANS, 7.5, SLATE)
    r = fp2.add_run(); set_run_font(r, MONO, 7.5, SLATE, rtl=False)
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), tag)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
        r._r.append(e)
    # title block
    add_text(doc, doc_code, size=8, color=ION_PRINT, mono=True, after=6, line=1.2)
    add_text(doc, title, size=22, bold=True, line=1.2, after=4, keep_next=True)
    if subtitle:
        add_text(doc, subtitle, size=11, color=SLATE, line=1.4, after=14)
    return doc

def h1(doc, text, num=None):
    t = (num + "  " + text) if num else text
    p = add_text(doc, t, size=14, bold=True, before=14, after=6, line=1.25, keep_next=True)
    return p

def h2(doc, text):
    return add_text(doc, text, size=11.5, bold=True, color=GRAPHITE, before=8, after=3, line=1.3, keep_next=True)

def para(doc, text, **kw):
    return add_text(doc, text, size=10.5, line=1.5, after=5, **kw)

def clause(doc, num, text):
    return add_runs(doc, [(num + "  ", True, ION_PRINT), (text, False)], size=10.5, line=1.5, after=5)

def bullets(doc, items, size=10.5):
    for it in items:
        add_runs(doc, [("•  ", True, ION_PRINT), (it, False)], size=size, line=1.45, after=3)

def eyebrow(doc, text):
    return add_text(doc, text, size=8, color=ION_PRINT, mono=True, after=2, line=1.2, keep_next=True)

def note(doc, text):
    t = doc.add_table(rows=1, cols=1); table_setup(t, [CONTENT_W])
    c = t.rows[0].cells[0]; shade(c, CARD); cell_borders(c, LINE, 0); cell_margins(c, 120, 120, 180, 180); clear_cell(c)
    add_text(c, text, size=9.5, color=GRAPHITE, line=1.45, after=0)
    spacer(doc, 6)
    return t

def fields_table(doc, rows, label_w=0.30, cols=1):
    """rows: list of (label, value). cols=2 puts two label/value pairs per row."""
    n = len(rows); per = cols
    nrows = (n + per - 1) // per
    t = doc.add_table(rows=nrows, cols=per * 2)
    pair_w = CONTENT_W / per
    widths = []
    for _ in range(per): widths += [int(pair_w * label_w), int(pair_w * (1 - label_w))]
    table_setup(t, widths)
    for i in range(nrows * per):
        lc = t.rows[i // per].cells[(i % per) * 2]; vc = t.rows[i // per].cells[(i % per) * 2 + 1]
        for c in (lc, vc):
            cell_borders(c, LINE, 6, ("bottom",)); cell_margins(c, 70, 70, 100, 100); clear_cell(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if i < n:
            label, value = rows[i]
            add_text(lc, label, size=9, color=SLATE, after=0, line=1.3)
            add_text(vc, value if value else " ", size=10.5, after=0, line=1.3)
        else:
            add_text(lc, " ", size=9, after=0); add_text(vc, " ", size=9, after=0)
        row_height(t.rows[i // per], 9)
    spacer(doc, 6)
    return t

def grid_table(doc, headers, rows, widths_frac, size=9.5, header_fill=GRAPHITE, min_row_mm=None, align=None, keep=None):
    t = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    widths = [int(CONTENT_W * f) for f in widths_frac]
    table_setup(t, widths)
    for j, h in enumerate(headers):
        c = t.rows[0].cells[j]; shade(c, header_fill); cell_borders(c, header_fill, 4); cell_margins(c); clear_cell(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        add_text(c, h, size=size - 0.5, color=PLATINUM, bold=True, after=0, line=1.25, align=align)
    for i, row in enumerate(rows):
        if isinstance(row, str):
            cells = t.rows[i + 1].cells
            for j, c in enumerate(cells):
                clear_cell(c); shade(c, CARD); cell_borders(c, LINE, 4); cell_margins(c)
                add_text(c, row if j == 1 else " ", size=size, bold=True, after=0, line=1.3)
            if min_row_mm: row_height(t.rows[i + 1], min_row_mm)
            continue
        for j, val in enumerate(row):
            c = t.rows[i + 1].cells[j]; cell_borders(c, LINE, 4); cell_margins(c); clear_cell(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i % 2 == 1: shade(c, "FAFBFC")
            bold = isinstance(val, tuple)
            txt = val[0] if bold else val
            add_text(c, txt if txt else " ", size=size, bold=bold, after=0, line=1.3, align=align)
        if min_row_mm: row_height(t.rows[i + 1], min_row_mm)
    if keep is None: keep = len(rows) <= 6
    if keep: keep_table(t)
    else:
        for row in t.rows:
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for c in t.rows[0].cells:
            for p in c.paragraphs: p.paragraph_format.keep_with_next = True
    spacer(doc, 6)
    return t

def checkbox_row(doc, items, per_row=4, size=9.5):
    rows = (len(items) + per_row - 1) // per_row
    t = doc.add_table(rows=rows, cols=per_row)
    table_setup(t, [int(CONTENT_W / per_row)] * per_row)
    for i in range(rows * per_row):
        c = t.rows[i // per_row].cells[i % per_row]; cell_borders(c, LINE, 0); cell_margins(c, 40, 40, 60, 60); clear_cell(c)
        if i < len(items):
            add_runs(c, [("☐  ", False, SLATE), (items[i], False)], size=size, after=0, line=1.3)
        else:
            add_text(c, " ", size=size, after=0)
    spacer(doc, 4)
    return t

def signature_block(doc, parties):
    """parties: list of (title, name_label). Renders side-by-side signature boxes."""
    n = len(parties)
    t = doc.add_table(rows=1, cols=n)
    table_setup(t, [int(CONTENT_W / n)] * n)
    for j, (title, sub) in enumerate(parties):
        c = t.rows[0].cells[j]; cell_borders(c, LINE, 0); cell_margins(c, 100, 100, 160, 160); clear_cell(c)
        add_text(c, title, size=10, bold=True, after=2, line=1.3)
        add_text(c, sub, size=9, color=SLATE, after=14, line=1.3)
        for lab in ("الاسم:", "التوقيع:", "التاريخ:"):
            add_runs(c, [(lab + "  ", False, SLATE), ("_______________________", False, CHROME)], size=9.5, after=6, line=1.3)
    keep_table(t)
    return t

def lines(doc, n=3, label=None, mm=9):
    if label: add_text(doc, label, size=9, color=SLATE, after=2, line=1.3, keep_next=True)
    t = doc.add_table(rows=n, cols=1); table_setup(t, [CONTENT_W])
    for i in range(n):
        c = t.rows[i].cells[0]; cell_borders(c, CHROME, 4, ("bottom",)); cell_margins(c, 40, 40, 60, 60); clear_cell(c)
        add_text(c, " ", size=9, after=0); row_height(t.rows[i], mm)
    spacer(doc, 6)

# ---------- schema-order fix so Word honours bidi/rtl ----------
_ORDER = {
    "pPr": ["pStyle","keepNext","keepLines","pageBreakBefore","framePr","widowControl","numPr","suppressLineNumbers","pBdr","shd","tabs","suppressAutoHyphens","kinsoku","wordWrap","overflowPunct","topLinePunct","autoSpaceDE","autoSpaceDN","bidi","adjustRightInd","snapToGrid","spacing","ind","contextualSpacing","mirrorIndents","suppressOverlap","jc","textDirection","textAlignment","textboxTightWrap","outlineLvl","divId","cnfStyle","rPr","sectPr","pPrChange"],
    "rPr": ["rStyle","rFonts","b","bCs","i","iCs","caps","smallCaps","strike","dstrike","outline","shadow","emboss","imprint","noProof","snapToGrid","vanish","webHidden","color","spacing","w","kern","position","sz","szCs","highlight","u","effect","bdr","shd","fitText","vertAlign","rtl","cs","em","lang","eastAsianLayout","specVanish","oMath"],
    "tblPr": ["tblStyle","tblpPr","tblOverlap","bidiVisual","tblStyleRowBandSize","tblStyleColBandSize","tblW","jc","tblCellSpacing","tblInd","tblBorders","shd","tblLayout","tblCellMar","tblLook","tblCaption","tblDescription"],
    "tcPr": ["cnfStyle","tcW","gridSpan","hMerge","vMerge","tcBorders","shd","noWrap","tcMar","textDirection","tcFitText","vAlign","hideMark"],
    "trPr": ["cnfStyle","divId","gridBefore","gridAfter","wBefore","wAfter","cantSplit","trHeight","tblHeader","tblCellSpacing","jc","hidden"],
}
def _fix_order(root):
    for tag, order in _ORDER.items():
        for el in root.iter(qn("w:" + tag)):
            kids = list(el)
            def key(k):
                name = k.tag.split("}")[-1]
                return order.index(name) if name in order else len(order)
            kids_sorted = sorted(kids, key=key)
            if kids_sorted != kids:
                for k in kids: el.remove(k)
                for k in kids_sorted: el.append(k)

def finalize(doc):
    _fix_order(doc.element.body)
    for s in doc.sections:
        for part in (s.header, s.footer):
            _fix_order(part._element)
    # document-level RTL defaults
    st = doc.styles["Normal"]
    pPr = st.element.get_or_add_pPr()
    b = OxmlElement("w:bidi"); pPr.insert(0, b)
    _fix_order(st.element)

import docx.document as _dd
_orig_save = _dd.Document.save
def _save(self, path):
    finalize(self); _orig_save(self, path)
_dd.Document.save = _save
